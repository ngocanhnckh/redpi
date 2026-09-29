#!/usr/bin/env python3
"""RedPi office files (original RedPi code): read Word, Excel, PowerPoint and PDF files as Markdown,
build .docx/.pptx from Markdown and .xlsx from CSV/JSON, merge/split PDFs, render pages to PNG.

Run through uv so nothing is installed globally, e.g.
  uv run --quiet --with python-docx python office.py read report.docx
Libraries: python-docx, openpyxl, python-pptx, pypdf (all open source).
"""
import csv, json, os, re, shutil, subprocess, sys, tempfile


def die(msg, code=1):
    print(f"office: {msg}", file=sys.stderr)
    sys.exit(code)


def md_cell(v):
    s = "" if v is None else str(v)
    return s.replace("|", "\\|").replace("\n", " ").strip()


def md_table(rows):
    rows = [list(r) for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(md_cell(c) for c in rows[0]) + " |", "| " + " | ".join("---" for _ in range(width)) + " |"]
    out += ["| " + " | ".join(md_cell(c) for c in r) + " |" for r in rows[1:]]
    return "\n".join(out)


# ---------- read ----------
def read_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = "".join(
                (f"**{r.text}**" if r.bold and r.text.strip() else f"*{r.text}*" if r.italic and r.text.strip() else r.text) for r in p.runs
            ).strip() or p.text.strip()
            if not text:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            m = re.match(r"Heading (\d)", style)
            if m:
                out.append("#" * int(m.group(1)) + " " + p.text.strip())
            elif style == "Title":
                out.append("# " + p.text.strip())
            elif "List" in style:
                m2 = re.search(r"(\d)$", style)
                level = int(m2.group(1)) - 1 if m2 else 0
                ppr = el.pPr
                if ppr is not None and ppr.numPr is not None and ppr.numPr.ilvl is not None:
                    level = int(ppr.numPr.ilvl.val)
                out.append("  " * level + ("1. " if "Number" in style else "- ") + text)
            else:
                out.append(text)
        elif tag == "tbl":
            t = Table(el, d)
            out.append(md_table([[c.text for c in row.cells] for row in t.rows]))
    sections = d.sections
    head = [p.text for s in sections for p in s.header.paragraphs if p.text.strip()] if sections else []
    if head:
        out.insert(0, "> Header: " + " / ".join(dict.fromkeys(head)))
    return "\n\n".join(x for x in out if x)


def read_xlsx(path, max_rows=200, max_cols=30):
    import openpyxl
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wf.worksheets:
        vs = wv[ws.title]
        out.append(f"## Sheet: {ws.title} ({ws.max_row} rows × {ws.max_column} columns{', hidden' if ws.sheet_state != 'visible' else ''})")
        rows = []
        for r in range(1, min(ws.max_row, max_rows) + 1):
            row = []
            for c in range(1, min(ws.max_column, max_cols) + 1):
                f, v = ws.cell(r, c).value, vs.cell(r, c).value
                if isinstance(f, str) and f.startswith("="):
                    row.append(f"{'' if v is None else v} `{f}`")
                else:
                    row.append(v)
            rows.append(row)
        table = md_table(rows)
        out.append(table or "(empty)")
        if ws.max_row > max_rows or ws.max_column > max_cols:
            out.append(f"(showing the first {min(ws.max_row, max_rows)} rows and {min(ws.max_column, max_cols)} columns)")
        if any(isinstance(ws.cell(r, c).value, str) and str(ws.cell(r, c).value).startswith("=") and vs.cell(r, c).value is None
               for r in range(1, min(ws.max_row, max_rows) + 1) for c in range(1, min(ws.max_column, max_cols) + 1)):
            out.append("(Some formulas have no cached value: the file was saved by a tool that does not calculate. Values show after opening in Excel/LibreOffice, or run `recalc`.)")
    return "\n\n".join(out)


def read_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        out.append(f"## Slide {i}{': ' + title if title else ''}")
        pics = 0
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                out.append(md_table([[c.text for c in row.cells] for row in shape.table.rows]))
            elif shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        out.append("  " * para.level + "- " + t)
            elif shape.shape_type == 13:
                pics += 1
        if pics:
            out.append(f"({pics} picture{'s' if pics > 1 else ''})")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("> Notes: " + slide.notes_slide.notes_text_frame.text.strip().replace("\n", " "))
    return "\n\n".join(out)


def read_pdf(path):
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True)
        if r.returncode == 0 and r.stdout.strip():
            pages = r.stdout.split("\f")
            return "\n\n".join(f"## Page {i}\n\n```\n{p.rstrip()}\n```" for i, p in enumerate(pages, 1) if p.strip())
    from pypdf import PdfReader
    rd = PdfReader(path)
    out = [f"## Page {i}\n\n{(p.extract_text() or '').strip() or '(no text: scanned image? render it and look)'}" for i, p in enumerate(rd.pages, 1)]
    fields = rd.get_fields() or {}
    if fields:
        out.append("## Form fields\n\n" + "\n".join(f"- {k}: {v.get('/V', '')}" for k, v in fields.items()))
    return "\n\n".join(out)


# ---------- write ----------
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def md_blocks(md):
    """(kind, payload) blocks: heading, bullet, number, table, code, para, rule."""
    lines = md.replace("\r\n", "\n").split("\n")
    i, blocks, para = 0, [], []
    flush = lambda: (blocks.append(("para", " ".join(para))), para.clear()) if para else None
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            blocks.append(("code", "\n".join(code))); i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", ln)
        if m:
            flush(); blocks.append(("heading", (len(m.group(1)), m.group(2).strip()))); i += 1; continue
        if re.match(r"^\s*(---|\*\*\*)\s*$", ln):
            flush(); blocks.append(("rule", None)); i += 1; continue
        m = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", ln)
        if m:
            flush(); blocks.append(("number" if m.group(2)[0].isdigit() else "bullet", (len(m.group(1)) // 2, m.group(3)))); i += 1; continue
        if ln.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            flush()
            rows = [[c.strip() for c in ln.strip().strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
            blocks.append(("table", rows)); continue
        if not ln.strip():
            flush()
        else:
            para.append(ln.strip())
        i += 1
    flush()
    return blocks


def docx_from_md(md, out):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    for kind, v in md_blocks(md):
        if kind == "heading":
            d.add_heading(v[1], level=0 if v[0] == 1 and not d.paragraphs else min(v[0], 9))
        elif kind in ("bullet", "number"):
            base = "List Bullet" if kind == "bullet" else "List Number"
            style = base if v[0] == 0 else f"{base} {min(v[0] + 1, 3)}"
            add_runs(d.add_paragraph(style=style), v[1])
        elif kind == "table":
            width = max(len(r) for r in v)
            t = d.add_table(rows=len(v), cols=width)
            t.style = "Table Grid"
            for r, row in enumerate(v):
                for c in range(width):
                    cell = t.cell(r, c)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        elif kind == "code":
            p = d.add_paragraph()
            r = p.add_run(v)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif kind == "rule":
            d.add_page_break()
        else:
            add_runs(d.add_paragraph(), v)
    d.save(out)


def pptx_from_md(md, out):
    """'# Title' + following text = title slide; each '## Heading' (or '---') starts a slide;
    bullets (indent for levels) and paragraphs fill it; 'Notes:' lines become speaker notes; a table becomes a table."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slides, cur = [], None
    for kind, v in md_blocks(md):
        if kind == "heading" and v[0] == 1 and not slides:
            cur = {"title": v[1], "items": [], "notes": [], "cover": True}; slides.append(cur); continue
        if kind == "heading" or kind == "rule":
            cur = {"title": v[1] if kind == "heading" else "", "items": [], "notes": [], "cover": False}; slides.append(cur); continue
        if cur is None:
            cur = {"title": "", "items": [], "notes": [], "cover": False}; slides.append(cur)
        if kind == "para" and v.lower().startswith("notes:"):
            cur["notes"].append(v[6:].strip())
        else:
            cur["items"].append((kind, v))
    for s in slides:
        if s["cover"]:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = s["title"]
            sub = " ".join(v if isinstance(v, str) else v[1] for k, v in s["items"] if k in ("para", "bullet", "number"))
            if sub and len(slide.placeholders) > 1:
                slide.placeholders[1].text = re.sub(r"[*`]", "", sub)
        else:
            tables = [v for k, v in s["items"] if k == "table"]
            slide = prs.slides.add_slide(prs.slide_layouts[5 if tables and len(tables) == len(s["items"]) else 1])
            slide.shapes.title.text = s["title"]
            texts = [(k, v) for k, v in s["items"] if k != "table"]
            if texts and len(slide.placeholders) > 1:
                tf = slide.placeholders[1].text_frame
                first = True
                for k, v in texts:
                    level, text = (v[0], v[1]) if k in ("bullet", "number") else (0, v)
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    p.level = min(level, 4)
                    p.text = re.sub(r"\*\*|`", "", text)
            for rows in tables:
                width = max(len(r) for r in rows)
                top = Inches(1.6 if not texts else 4.2)
                shape = slide.shapes.add_table(len(rows), width, Inches(0.6), top, prs.slide_width - Inches(1.2), Inches(0.4) * len(rows))
                for r, row in enumerate(rows):
                    for c in range(width):
                        cell = shape.table.cell(r, c)
                        cell.text = re.sub(r"\*\*|`", "", row[c] if c < len(row) else "")
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(14)
        if s["notes"]:
            slide.notes_slide.notes_text_frame.text = "\n".join(s["notes"])
    prs.save(out)


def xlsx_from_table(src, out, sheet="Sheet1"):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    if src.lower().endswith(".json"):
        data = json.load(open(src, encoding="utf-8"))
        if isinstance(data, dict):
            data = data.get("rows") or data.get("data") or [data]
        header = list(dict.fromkeys(k for row in data for k in row.keys()))
        rows = [header] + [[row.get(k) for k in header] for row in data]
    else:
        with open(src, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))

    def typed(v):
        if not isinstance(v, str):
            return v
        s = v.strip()
        if re.fullmatch(r"-?\d+", s) and not (len(s) > 1 and s.startswith("0")):
            return int(s)
        if re.fullmatch(r"-?\d*\.\d+", s):
            return float(s)
        return v

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v if r == 1 else typed(v))
    for c in range(1, len(rows[0]) + 1 if rows else 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        width = max((len(str(ws.cell(r, c).value or "")) for r in range(1, min(len(rows), 500) + 1)), default=8)
        ws.column_dimensions[get_column_letter(c)].width = min(60, max(8, width + 2))
    if rows:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(out)


def pdf_merge(inputs, out):
    from pypdf import PdfWriter
    w = PdfWriter()
    for p in inputs:
        w.append(p)
    with open(out, "wb") as f:
        w.write(f)


def page_list(spec, n):
    pages = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            pages += range(int(a or 1), int(b or n) + 1)
        elif part:
            pages.append(int(part))
    bad = [p for p in pages if p < 1 or p > n]
    if bad:
        die(f"pages {bad} are outside 1-{n}")
    return pages


def pdf_pages(path, spec, out):
    from pypdf import PdfReader, PdfWriter
    rd = PdfReader(path)
    w = PdfWriter()
    for p in page_list(spec, len(rd.pages)):
        w.add_page(rd.pages[p - 1])
    with open(out, "wb") as f:
        w.write(f)


def soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def to_pdf(path, outdir):
    exe = soffice()
    if not exe:
        die("converting Office files needs LibreOffice (install it: sudo apt install libreoffice-core, or brew install --cask libreoffice)")
    r = subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", outdir, path], capture_output=True, text=True, timeout=300)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    if r.returncode != 0 or not os.path.exists(pdf):
        die(f"LibreOffice could not convert {path}: {(r.stderr or r.stdout).strip()[:400]}")
    return pdf


def render(path, outdir, pages=None, dpi=80):
    os.makedirs(outdir, exist_ok=True)
    if not shutil.which("pdftoppm"):
        die("rendering needs pdftoppm (poppler-utils: sudo apt install poppler-utils, or brew install poppler)")
    pdf = path if path.lower().endswith(".pdf") else to_pdf(path, tempfile.mkdtemp())
    base = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0])
    args = ["pdftoppm", "-png", "-r", str(dpi)]
    if pages:
        a, _, b = pages.partition("-")
        args += ["-f", a, "-l", b or a]
    subprocess.run(args + [pdf, base], check=True)
    made = sorted(f for f in os.listdir(outdir) if f.startswith(os.path.basename(base)) and f.endswith(".png"))
    return [os.path.join(outdir, f) for f in made]


def recalc(path):
    """Recalculate formulas and cache their values by round-tripping through LibreOffice."""
    exe = soffice()
    if not exe:
        die("recalculating needs LibreOffice (install it: sudo apt install libreoffice-core, or brew install --cask libreoffice)")
    tmp = tempfile.mkdtemp()
    r = subprocess.run([exe, "--headless", "--convert-to", "xlsx", "--outdir", tmp, path], capture_output=True, text=True, timeout=300)
    out = os.path.join(tmp, os.path.basename(path))
    if r.returncode != 0 or not os.path.exists(out):
        die(f"LibreOffice could not recalculate {path}")
    shutil.copyfile(out, path)


READERS = {".docx": read_docx, ".xlsx": read_xlsx, ".xlsm": read_xlsx, ".pptx": read_pptx, ".pdf": read_pdf}

USAGE = """usage: office.py <command> ...
  read <file> [--max CHARS]              Word/Excel/PowerPoint/PDF as Markdown
  docx <in.md> <out.docx>                Markdown -> Word
  pptx <in.md> <out.pptx>                Markdown outline -> PowerPoint (## starts a slide)
  xlsx <in.csv|in.json> <out.xlsx> [--sheet NAME]
  merge <out.pdf> <in1.pdf> <in2.pdf>...
  pages <in.pdf> <spec e.g. 1-3,5> <out.pdf>
  render <file> <outdir> [--pages 1-3] [--dpi 80]   pages as PNG (Office files need LibreOffice)
  topdf <file> <outdir>                  Office file -> PDF (needs LibreOffice)
  recalc <file.xlsx>                     recalculate formulas (needs LibreOffice)"""


def main(argv):
    if len(argv) < 2 or argv[1] in ("-h", "--help", "help"):
        print(USAGE); return
    cmd, a = argv[1], argv[2:]
    opt = lambda name, default=None: a[a.index(name) + 1] if name in a and a.index(name) + 1 < len(a) else default
    if cmd == "read":
        ext = os.path.splitext(a[0])[1].lower()
        if ext not in READERS:
            die(f"cannot read {ext or 'that'} files (docx, xlsx, pptx, pdf)")
        text = READERS[ext](a[0])
        mx = int(opt("--max", "60000"))
        print(text if len(text) <= mx else text[:mx] + f"\n\n…(cut at {mx} characters of {len(text)}; use --max for more)")
    elif cmd == "docx":
        docx_from_md(open(a[0], encoding="utf-8").read(), a[1]); print(f"wrote {a[1]}")
    elif cmd == "pptx":
        pptx_from_md(open(a[0], encoding="utf-8").read(), a[1]); print(f"wrote {a[1]}")
    elif cmd == "xlsx":
        xlsx_from_table(a[0], a[1], opt("--sheet", "Sheet1")); print(f"wrote {a[1]}")
    elif cmd == "merge":
        pdf_merge(a[1:], a[0]); print(f"wrote {a[0]}")
    elif cmd == "pages":
        pdf_pages(a[0], a[1], a[2]); print(f"wrote {a[2]}")
    elif cmd == "render":
        for p in render(a[0], a[1], opt("--pages"), int(opt("--dpi", "80"))):
            print(p)
    elif cmd == "topdf":
        print(to_pdf(a[0], a[1]))
    elif cmd == "recalc":
        recalc(a[0]); print(f"recalculated {a[0]}")
    else:
        die(USAGE)


if __name__ == "__main__":
    main(sys.argv)
